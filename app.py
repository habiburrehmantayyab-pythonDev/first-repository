
import re
import io
import streamlit as st
from pypdf import PdfReader
from docx import Document

st.set_page_config(
    page_title="ATS Resume Checker",
    page_icon="📄",
    layout="wide"
)

st.title("ATS Resume Checker")
st.write(
    "Upload your resume, enter a job description, "
    "and analyze your resume out of 100."
)

st.info(
    "This is an educational, rule-based ATS checker. "
    "Its score is an estimate, not a real employer ATS score."
)


def extract_text(uploaded_file):
    """Extract text from PDF or DOCX files."""

    file_bytes = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif filename.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))

        paragraphs = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(
                    " ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs)

    else:
        raise ValueError("Please upload a PDF or DOCX file.")


def normalize(text):
    return re.sub(r"\s+", " ", text.lower()).strip()


def analyze_resume(resume_text, job_description):

    text = normalize(resume_text)
    job_text = normalize(job_description)

    score = 0
    strengths = []
    missing = []
    suggestions = []
    breakdown = {}

    # 1. Contact Details: 15 Points
    contact_score = 0

    if re.search(
        r"\b[\w.+-]+@[\w.-]+\.[a-z]{2,}\b", text
    ):
        contact_score += 8
        strengths.append("Email address found.")
    else:
        missing.append("Email address")
        suggestions.append(
            "Add a professional email address."
        )

    if re.search(r"(\+?\d[\d\s().-]{7,}\d)", resume_text):
        contact_score += 7
        strengths.append("Possible phone number found.")
    else:
        missing.append("Phone number")
        suggestions.append(
            "Add your phone number."
        )

    breakdown["Contact Details"] = (contact_score, 15)
    score += contact_score

    # 2. Resume Sections: 25 Points
    section_groups = {
        "Summary / Objective": [
            "summary", "objective", "professional profile"
        ],
        "Education": [
            "education", "academic background", "qualification"
        ],
        "Experience": [
            "experience", "employment history", "work history"
        ],
        "Skills": [
            "skills", "technical skills", "core competencies"
        ],
        "Projects": [
            "projects", "personal projects", "academic projects"
        ],
    }

    section_score = 0

    for section, keywords in section_groups.items():
        if any(word in text for word in keywords):
            section_score += 5
            strengths.append(f"{section} section found.")
        else:
            missing.append(section)
            suggestions.append(
                f"Consider adding a relevant {section} section."
            )

    breakdown["Resume Sections"] = (section_score, 25)
    score += section_score

    # 3. Job Description Keyword Matching: 40 Points
    stop_words = {
        "the", "and", "for", "with", "you", "your", "are",
        "our", "this", "that", "will", "have", "has", "from",
        "into", "who", "what", "when", "where", "their", "they",
        "job", "role", "work", "working", "years", "year",
        "ability", "skills", "experience", "required", "preferred",
        "using", "use", "including", "such", "must", "can"
    }

    job_words = set(
        re.findall(r"[a-zA-Z][a-zA-Z+#.-]*", job_text)
    )

    resume_words = set(
        re.findall(r"[a-zA-Z][a-zA-Z+#.-]*", text)
    )

    important_words = {
        word for word in job_words
        if word not in stop_words and len(word) > 2
    }

    matched_words = important_words & resume_words

    if important_words:
        match_ratio = len(matched_words) / len(important_words)
        keyword_score = round(match_ratio * 40)
    else:
        keyword_score = 0

    breakdown["Job Keyword Matching"] = (keyword_score, 40)
    score += keyword_score

    if matched_words:
        strengths.append(
            "Matching keywords: "
            + ", ".join(sorted(matched_words)[:15])
        )

    unmatched_words = important_words - resume_words

    if unmatched_words:
        missing.append(
            "Some job description keywords are not found."
        )

        suggestions.append(
            "Review these possible missing keywords: "
            + ", ".join(sorted(unmatched_words)[:20])
        )

    # 4. Content Quality: 20 Points
    quality_score = 0

    word_count = len(text.split())

    if word_count >= 150:
        quality_score += 5
        strengths.append(
            "Resume contains a reasonable amount of text."
        )
    else:
        missing.append("More resume detail may be needed.")
        suggestions.append(
            "Add relevant details about your work, skills, "
            "education, and projects."
        )

    if re.search(
        r"\b(19|20)\d{2}\b|"
        r"\b(january|february|march|april|may|june|july|"
        r"august|september|october|november|december)\b",
        text
    ):
        quality_score += 5
        strengths.append("Possible dates or timelines found.")
    else:
        suggestions.append(
            "Add dates to education and work experience where relevant."
        )

    action_verbs = [
        "developed", "managed", "created", "designed",
        "implemented", "improved", "built", "led",
        "analyzed", "automated", "organized", "delivered"
    ]

    found_verbs = [
        verb for verb in action_verbs if verb in text
    ]

    if found_verbs:
        quality_score += 5
        strengths.append("Action-oriented wording found.")
    else:
        suggestions.append(
            "Use strong action verbs such as developed, "
            "managed, designed, or implemented."
        )

    if re.search(
        r"\b\d+%|\b\d+\+?\s*(users|clients|projects|teams)\b",
        text
    ):
        quality_score += 5
        strengths.append("Possible measurable achievements found.")
    else:
        suggestions.append(
            "Where truthful, quantify achievements using "
            "numbers, percentages, or project outcomes."
        )

    breakdown["Content Quality"] = (quality_score, 20)
    score += quality_score

    score = min(score, 100)

    return {
        "score": score,
        "strengths": strengths,
        "missing": missing,
        "suggestions": list(dict.fromkeys(suggestions)),
        "breakdown": breakdown,
        "matched": sorted(matched_words),
        "unmatched": sorted(unmatched_words),
        "word_count": word_count,
    }


# Resume Upload
uploaded_file = st.file_uploader(
    "Upload Your Resume",
    type=["pdf", "docx"]
)

# Job Description Input
job_description = st.text_area(
    "Paste the Job Description",
    height=220,
    placeholder=(
        "Example: We are looking for a Python developer "
        "with Python, SQL, Git, and problem-solving skills."
    )
)

# Analyze Button
if st.button("Analyze Resume", type="primary"):

    if uploaded_file is None:
        st.error("Please upload a PDF or DOCX resume.")

    elif not job_description.strip():
        st.error("Please enter the job description.")

    else:
        try:
            resume_text = extract_text(uploaded_file)

            if not resume_text.strip():
                st.error(
                    "No readable text found. The PDF may be scanned. "
                    "Try a text-based PDF or DOCX file."
                )

            else:
                result = analyze_resume(
                    resume_text,
                    job_description
                )

                st.divider()
                st.header("Resume Analysis Report")

                col1, col2 = st.columns(2)

                with col1:
                    st.metric(
                        "Estimated ATS Score",
                        f"{result['score']}/100"
                    )

                    st.progress(result["score"] / 100)

                with col2:
                    st.metric(
                        "Resume Word Count",
                        result["word_count"]
                    )

                if result["score"] >= 80:
                    st.success(
                        "Good estimated match. Review suggestions "
                        "before applying."
                    )

                elif result["score"] >= 60:
                    st.warning(
                        "Moderate estimated match. "
                        "Some improvements may help."
                    )

                else:
                    st.error(
                        "Several improvements may be needed. "
                        "Review the report below."
                    )

                # Score Breakdown
                st.subheader("Score Breakdown")

                for category, values in result["breakdown"].items():
                    earned, maximum = values

                    st.write(
                        f"**{category}:** {earned}/{maximum}"
                    )

                    st.progress(earned / maximum)

                # Strengths and Missing Information
                left, right = st.columns(2)

                with left:
                    st.subheader("What's Right")

                    if result["strengths"]:
                        for item in result["strengths"]:
                            st.write(f"- {item}")
                    else:
                        st.write(
                            "No strengths detected by these rules."
                        )

                with right:
                    st.subheader("Missing / Needs Review")

                    if result["missing"]:
                        for item in result["missing"]:
                            st.write(f"- {item}")
                    else:
                        st.write(
                            "No missing items detected by these rules."
                        )

                # Suggestions
                st.subheader("Improvement Suggestions")

                if result["suggestions"]:
                    for item in result["suggestions"]:
                        st.write(f"- {item}")
                else:
                    st.write(
                        "No additional suggestions were generated."
                    )

                # Matched Keywords
                with st.expander("View Matched Job Keywords"):
                    st.write(
                        ", ".join(result["matched"])
                        or "No matches found."
                    )

                # Missing Keywords
                with st.expander("View Possible Missing Keywords"):
                    st.write(
                        ", ".join(result["unmatched"])
                        or "No unmatched keywords found."
                    )

                # Build Downloadable Report
                report = [
                    "ATS RESUME CHECKER REPORT",
                    f"Estimated Score: {result['score']}/100",
                    f"Word Count: {result['word_count']}",
                    "",
                    "SCORE BREAKDOWN:"
                ]

                for category, values in result["breakdown"].items():
                    report.append(
                        f"{category}: {values[0]}/{values[1]}"
                    )

                report.extend(["", "STRENGTHS:"])
                report.extend(
                    f"- {item}" for item in result["strengths"]
                )

                report.extend(["", "MISSING / NEEDS REVIEW:"])
                report.extend(
                    f"- {item}" for item in result["missing"]
                )

                report.extend(["", "SUGGESTIONS:"])
                report.extend(
                    f"- {item}" for item in result["suggestions"]
                )

                st.download_button(
                    "Download Analysis Report",
                    data="\n".join(report),
                    file_name="ats_resume_report.txt",
                    mime="text/plain"
                )

        except Exception:
            st.error(
                "Could not analyze this file. Please check that "
                "it is a valid PDF or DOCX file."
            )
            
                    
